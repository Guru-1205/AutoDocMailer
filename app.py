from flask import Flask, render_template, request, send_file
from io import BytesIO
from docx import Document
from docx.shared import Pt
from docx.oxml import OxmlElement
from docx.oxml.ns import nsdecls

app = Flask(__name__)

template_path = 'static/template.docx'
custom_font_path = 'static/fonts/times new roman bold.ttf'


@app.route('/')
def index():
    return render_template('index.html')


@app.route('/process', methods=['POST'])
def process():
    # Get user input from the form
    replace_dict = {
        # '{{NAME}}': request.form['name'],
        # '{{DATE}}': request.form['date'],
        # '{{ADDRESS}}': request.form['address']
        '{{activityName}}': request.form['name'],
        '{{activityFromDate}}': request.form['fromDate'],
        '{{activityYear}}': request.form['year'],
        '{{activityDate}}': request.form['date'],
        '{{activityPlace}}': request.form['place'],
        '{{activityDuration}}': request.form['timing'],
        '{{activityDay}}': request.form['day']
    }

    # Load the Word document template
    with open(template_path, 'rb') as doc_stream:
        doc = Document(BytesIO(doc_stream.read()))

    # Replace placeholders with user input
    for paragraph in doc.paragraphs:
        for key, value in replace_dict.items():
            if key in paragraph.text:
                paragraph.text = paragraph.text.replace(key, str(value))

    # Create an in-memory stream for the modified document
    output_stream = BytesIO()
    doc.save(output_stream)

    # Move the stream cursor to the beginning
    output_stream.seek(0)

    # Serve the modified document directly without saving it
    return send_file(output_stream,
                     mimetype='application/vnd.openxmlformats-officedocument.wordprocessingml.document',
                     as_attachment=True, download_name='output_document.docx')


if __name__ == '__main__':
    app.run(debug=True)
